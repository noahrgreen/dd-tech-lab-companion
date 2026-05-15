#!/usr/bin/env python3
from pathlib import Path
import json
import csv
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
ARTICLES = ROOT / 'articles'
FOUNDATION = Path('/Users/noahgreen/NGO/outbox/owner_review/dd_tech_lab_kickoff_2026-05-10')
SEEDS_FOUNDRY = FOUNDATION / 'palantir_foundry' / 'seeds'
SEEDS_EXCEL = FOUNDATION / 'excel_fraud' / 'seeds'
INDEX = ROOT / 'articles_index.md'


def write(path: Path, text: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text)


def dump_json(path: Path, data):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2) + '\n')


def dump_csv(path: Path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('w', newline='') as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def build_workbook(path: Path, title: str, purpose: str, tabs):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = 'README'
    ws['A1'] = title
    ws['A2'] = purpose
    ws['A4'] = 'This is a starter workbook scaffold. Populate formulas and worked examples once the article draft is finalized.'
    for tab in tabs:
        ws2 = wb.create_sheet(tab)
        ws2['A1'] = tab
        ws2['A2'] = 'Starter sheet scaffold'
    wb.save(path)


FOUNDRY_SPECS = [
    {
        'dir': '003_workshop_application_patterns_investigators',
        'title': 'Workshop Application Patterns for Counterparty Risk Investigators — Single-Pane-of-Glass Investigator Workflows',
        'article_no': '003',
        'seed': '003_workshop-application-patterns-investigators_seed_DRAFT_v1_2026-05-11.md',
        'files': {
            'README.md': '# Foundry Article 003 — Starter Companion Bundle\n\nThis starter bundle captures the artifact shape promised by the seed: investigator-facing Workshop layout, action definitions, and a minimal synthetic case payload.\n',
            'workshop_module/workshop_layout.yaml': 'module: counterparty_risk_investigation\nlayout: split_pane\nprimary_sections:\n  - header\n  - summary\n  - relationships\nsecondary_sections:\n  - transaction_volume_time_series_chart\n  - ownership_graph_visualization\n  - sanctions_list_full_text\n',
            'action_templates/investigator_actions.yaml': 'actions:\n  - name: ElevateRiskRating\n    required_fields: [justification, supporting_evidence_object_set]\n  - name: InitiateEnhancedDueDiligence\n    required_fields: [reason, linked_counterparty]\n',
            'synthetic_data/investigation_case_minimal.json': json.dumps({'counterparty_id': 'CP-SYN-003', 'risk_rating': 'high', 'sanctions_hits': 1, 'ubo_count': 2}, indent=2),
        },
    },
    {
        'dir': '004_aip_adverse_media_summarization',
        'title': 'AIP-Driven Adverse-Media Summarization for DD Engagements — Prompt Patterns, Grounding Strategy, Hallucination Controls',
        'article_no': '004',
        'seed': '004_aip-adverse-media-summarization_seed_DRAFT_v1_2026-05-11.md',
        'files': {
            'README.md': '# Foundry Article 004 — Starter Companion Bundle\n\nThis starter bundle captures the prompt, output schema, and hallucination-control gates described in the seed.\n',
            'prompts/adverse_media_summarization_prompt.md': '# System prompt\n\nSummarize only grounded adverse-media facts. Every claim requires an attached citation.\n',
            'schemas/grounded_summary_schema.json': json.dumps({'summary': 'string', 'citations': [{'source_id': 'string', 'claim': 'string'}], 'risk_flags': ['string']}, indent=2),
            'evaluations/hallucination_gate.yaml': 'gates:\n  - all factual assertions cite a source\n  - no unsupported jurisdiction claims\n  - no risk escalation without evidence\n',
        },
    },
    {
        'dir': '005_actions_framework_audit_trail_discipline',
        'title': 'Foundry Actions Framework for Audit-Trail Discipline — Risk-Rating Changes, Regulatory Documentation, and the Examiner-Ready Audit Log',
        'article_no': '005',
        'seed': '005_actions-framework-audit-trail-discipline_seed_DRAFT_v1_2026-05-11.md',
        'files': {
            'README.md': '# Foundry Article 005 — Starter Companion Bundle\n\nThis starter bundle defines the action payload, audit-log schema, and retention matrix implied by the seed.\n',
            'action_templates/risk_rating_change_action.yaml': 'action: ElevateRiskRating\nfields:\n  - counterparty_id\n  - prior_rating\n  - new_rating\n  - justification\n  - supporting_evidence_object_set\n',
            'audit_log/risk_decision_audit_log_schema.json': json.dumps({'action_id': 'string', 'counterparty_id': 'string', 'actor': 'string', 'timestamp': 'iso8601', 'old_value': 'string', 'new_value': 'string', 'evidence_ids': ['string']}, indent=2),
            'retention/retention_and_approvals_matrix.yaml': 'records:\n  risk_decision_log: 7_years\n  supporting_evidence_snapshot: 7_years\napprovals:\n  critical_rating_change: compliance_manager\n',
        },
    },
    {
        'dir': '006_code_repositories_python_pyspark',
        'title': 'Code Repositories in Foundry — When to Embed Python / PySpark in the Pipeline (and When to Stay in Pipeline Builder)',
        'article_no': '006',
        'seed': '006_code-repositories-python-pyspark_seed_DRAFT_v1_2026-05-11.md',
        'files': {
            'README.md': '# Foundry Article 006 — Starter Companion Bundle\n\nThis starter bundle frames the decision boundary between Pipeline Builder and Code Repositories with a concrete transform stub.\n',
            'transforms/pyspark_silver_to_gold.py': 'from pyspark.sql import DataFrame\n\n\ndef build_gold_counterparty(df: DataFrame) -> DataFrame:\n    return df\n',
            'tests/transform_contract_tests.md': '# Contract tests\n\n- idempotent output under repeat runs\n- schema drift raises explicit mapping error\n- audit metadata columns always present\n',
        },
        'csvs': {
            'decision_matrix/pipeline_vs_code_repo_matrix.csv': (
                ['criterion', 'pipeline_builder', 'code_repo'],
                [
                    ['simple schema mapping', 'preferred', 'not needed'],
                    ['multi-source custom matching', 'limited', 'preferred'],
                    ['PySpark dependency required', 'unsupported', 'required'],
                ],
            ),
        },
    },
    {
        'dir': '007_quiver_ad_hoc_counterparty_queries',
        'title': 'Quiver for Ad-Hoc Counterparty Queries — Lightweight Investigator Tooling Without a Full Workshop Build',
        'article_no': '007',
        'seed': '007_quiver-ad-hoc-counterparty-queries_seed_DRAFT_v1_2026-05-11.md',
        'files': {
            'README.md': '# Foundry Article 007 — Starter Companion Bundle\n\nThis starter bundle provides a lightweight query pack and small synthetic query examples.\n',
            'query_pack/quiver_counterparty_queries.yaml': 'queries:\n  - name: exposure_to_issuer\n    prompt: Which counterparties have direct or indirect exposure to issuer X?\n  - name: sanctions_overlap\n    prompt: Which counterparties share a beneficial owner with any active sanctions hit?\n',
        },
        'csvs': {
            'synthetic_data/counterparty_query_examples.csv': (
                ['query_name', 'input', 'expected_shape'],
                [
                    ['exposure_to_issuer', 'issuer_name', 'counterparty list with exposure path'],
                    ['sanctions_overlap', 'sanctions_entity', 'counterparties + shared owner'],
                ],
            ),
        },
    },
    {
        'dir': '008_time_series_counterparty_risk_trajectories',
        'title': 'Time Series in Foundry — Modeling Counterparty Risk Trajectories and Early-Warning Indicator Pipelines',
        'article_no': '008',
        'seed': '008_time-series-counterparty-risk-trajectories_seed_DRAFT_v1_2026-05-11.md',
        'files': {
            'README.md': '# Foundry Article 008 — Starter Companion Bundle\n\nThis starter bundle seeds the feature schema, synthetic trajectory data, and alerting pipeline shape described in the seed.\n',
            'feature_schema/risk_trajectory_features.yaml': 'features:\n  - monthly_txn_volume_change_pct\n  - sanctions_screening_hits_rolling_90d\n  - adverse_media_mentions_rolling_30d\n  - risk_rating_change_count_rolling_180d\n',
            'pipeline_templates/trajectory_alert_pipeline.yaml': 'inputs:\n  - counterparty_monthly_features\ntransforms:\n  - compute_zscores\n  - threshold_breach_flags\n  - emit_alert_objects\n',
        },
        'csvs': {
            'synthetic_data/counterparty_risk_timeseries.csv': (
                ['counterparty_id', 'month', 'risk_score', 'adverse_media_count', 'sanctions_hits'],
                [
                    ['CP-SYN-008', '2026-01', '42', '0', '0'],
                    ['CP-SYN-008', '2026-02', '47', '1', '0'],
                    ['CP-SYN-008', '2026-03', '61', '2', '1'],
                ],
            ),
        },
    },
    {
        'dir': '009_beneficial_ownership_networks',
        'title': 'Foundry Ontology Design for Beneficial-Ownership Networks — Person, Entity, Officer, and Jurisdiction Modeling at Bank-Group Scale',
        'article_no': '009',
        'seed': '009_ontology-beneficial-ownership-networks_seed_DRAFT_v1_2026-05-11.md',
        'files': {
            'README.md': '# Foundry Article 009 — Starter Companion Bundle\n\nThis starter bundle defines the ontology, synthetic UBO entities/links, and a query starter for cross-jurisdiction ownership traversal.\n',
            'ontology_schema/beneficial_ownership_ontology.yaml': 'object_types:\n  Person:\n    primary_key: person_id\n  Entity:\n    primary_key: entity_id\n  Jurisdiction:\n    primary_key: country_code\nlink_types:\n  OWNS:\n    properties: [ownership_percentage, effective_from, effective_to, disclosure_source]\n  IS_OFFICER_OF:\n    properties: [role_title, effective_from, effective_to]\n',
            'queries/beneficial_ownership_path_queries.cypher': 'MATCH p=(person:Person)-[:OWNS*1..6]->(entity:Entity)\nRETURN person, entity, p\nLIMIT 25;\n',
        },
        'csvs': {
            'synthetic_data/ubo_entities.csv': (
                ['entity_id', 'legal_name', 'jurisdiction'],
                [
                    ['E-001', 'North Harbor Trading Ltd', 'KY'],
                    ['E-002', 'North Harbor Operating LLC', 'US'],
                ],
            ),
            'synthetic_data/ubo_links.csv': (
                ['source_id', 'target_id', 'ownership_pct', 'effective_from'],
                [
                    ['P-001', 'E-001', '60', '2024-01-01'],
                    ['E-001', 'E-002', '100', '2024-01-01'],
                ],
            ),
        },
    },
    {
        'dir': '010_platform_tradeoff_framework',
        'title': 'Foundry vs Snowflake + DBT vs Databricks for DD Analytics — Architecture Trade-Off Framework',
        'article_no': '010',
        'seed': '010_foundry-vs-snowflake-dbt-databricks_seed_DRAFT_v1_2026-05-11.md',
        'files': {
            'README.md': '# Foundry Article 010 — Starter Companion Bundle\n\nThis starter bundle holds the comparison scorecard and weighting template for the platform-tradeoff article.\n',
            'templates/platform_evaluation_weights.yaml': 'weights:\n  ontology_layer: 0.20\n  investigator_ui: 0.15\n  pipeline_transforms: 0.15\n  llm_integration: 0.10\n  audit_trail: 0.15\n  cost_model: 0.10\n  skill_alignment: 0.10\n  vendor_lock_in: 0.05\n',
        },
        'csvs': {
            'decision_matrix/platform_tradeoff_scorecard.csv': (
                ['dimension', 'foundry', 'snowflake_dbt', 'databricks'],
                [
                    ['ontology_layer', 'built_in', 'build_yourself', 'partial'],
                    ['investigator_ui', 'Workshop', 'custom_app', 'Lakeview_partial'],
                    ['audit_trail', 'Actions', 'custom', 'custom'],
                ],
            ),
        },
    },
]

EXCEL_SPECS = [
    {
        'dir': 'excel_002_same_same_different',
        'title': 'Same-Same-Different in Excel',
        'article_no': '002',
        'seed': '002_same-same-different-in-excel_seed_DRAFT_v1_2026-05-13.md',
        'workbook_path': ROOT / 'workbooks/002-same-same-different.xlsx',
        'tabs': ['Inputs', 'Pairing Logic', 'Exception Review', 'Checks'],
    },
    {
        'dir': 'excel_003_linear_regression_outlier_detection',
        'title': 'Linear Regression Outlier Detection in Excel',
        'article_no': '003',
        'seed': '003_linear-regression-outlier-detection-in-excel_seed_DRAFT_v1_2026-05-13.md',
        'workbook_path': ROOT / 'workbooks/003-linear-regression-outlier-detection.xlsx',
        'tabs': ['Inputs', 'Regression Model', 'Residual Review', 'Checks'],
    },
    {
        'dir': 'excel_004_geospatial_address_mapping',
        'title': 'Geospatial Address Mapping in Excel',
        'article_no': '004',
        'seed': '004_geospatial-address-mapping-in-excel_seed_DRAFT_v1_2026-05-13.md',
        'workbook_path': ROOT / 'workbooks/004-geospatial-address-mapping.xlsx',
        'tabs': ['Inputs', 'Coordinate Prep', 'Distance Review', 'Checks'],
    },
    {
        'dir': 'excel_007_round_number_bias_threshold_avoidance',
        'title': 'Round-Number Bias and Threshold Avoidance in Excel',
        'article_no': '007',
        'seed': '007_round-number-bias-threshold-avoidance_seed_DRAFT_v1_2026-05-13.md',
        'workbook_path': ROOT / 'workbooks/007-round-number-bias-threshold-avoidance.xlsx',
        'tabs': ['Inputs', 'Digit Flags', 'Threshold Bands', 'Checks'],
    },
    {
        'dir': 'excel_010_correlation_diagnostics_journal_entry_pairs',
        'title': 'Correlation Diagnostics for Journal-Entry Pairs in Excel',
        'article_no': '010',
        'seed': '010_correlation-diagnostics-journal-entry-pairs_seed_DRAFT_v1_2026-05-13.md',
        'workbook_path': ROOT / 'workbooks/010-correlation-diagnostics-journal-entry-pairs.xlsx',
        'tabs': ['Inputs', 'Pairing Matrix', 'Correlation Signals', 'Checks'],
    },
]


def build_foundry():
    created = []
    for spec in FOUNDRY_SPECS:
        base = ARTICLES / spec['dir']
        base.mkdir(parents=True, exist_ok=True)
        manifest = {
            'article_no': spec['article_no'],
            'title': spec['title'],
            'series': 'Palantir Foundry for Due Diligence',
            'status': 'starter_bundle',
            'source_seed': str(SEEDS_FOUNDRY / spec['seed']),
            'quality_standard': 'Parseable starter artifacts aligned to the article seed; replace stubs with production-grade companions once article draft is finalized.',
        }
        dump_json(base / 'artifact_manifest.json', manifest)
        for rel, content in spec['files'].items():
            write(base / rel, content)
        for rel, (header, rows) in spec.get('csvs', {}).items():
            dump_csv(base / rel, header, rows)
        created.append(spec['dir'])
    return created


def build_excel():
    created = []
    for spec in EXCEL_SPECS:
        base = ARTICLES / spec['dir']
        base.mkdir(parents=True, exist_ok=True)
        manifest = {
            'article_no': spec['article_no'],
            'title': spec['title'],
            'series': 'Excel Fraud Analytics',
            'status': 'starter_bundle',
            'source_seed': str(SEEDS_EXCEL / spec['seed']),
            'workbook_path': str(spec['workbook_path'].relative_to(ROOT)),
            'quality_standard': 'Workbook starter scaffold only; formulas, examples, and assertions must be finalized against the canonical article before public linking.',
        }
        dump_json(base / 'artifact_manifest.json', manifest)
        write(base / 'README.md', f"# Excel Article {spec['article_no']} — Starter Companion Bundle\n\nStarter workbook scaffold for **{spec['title']}**. The workbook exists now so the bundle path is stable; analytical formulas and worked examples remain to be finalized against the canonical article draft.\n")
        build_workbook(spec['workbook_path'], spec['title'], 'Starter workbook scaffold aligned to the article seed.', spec['tabs'])
        created.append(spec['dir'])
    return created


def update_index():
    lines = [
        '# DD Tech Lab — Article-by-Article Artifact Index',
        '',
        '| Article | Title | Artifacts |',
        '|--------|-------|-----------|',
        '| **Foundry 001** | Foundry Ontology Design for Counterparty Risk Investigations | [ontology_schema](articles/001_foundry_ontology_counterparty_risk/ontology_schema/) · [synthetic_data](articles/001_foundry_ontology_counterparty_risk/synthetic_data/) · [foundry_sdk_stubs](articles/001_foundry_ontology_counterparty_risk/foundry_sdk_stubs/) · [workshop_module](articles/001_foundry_ontology_counterparty_risk/workshop_module/) · [action_templates](articles/001_foundry_ontology_counterparty_risk/action_templates/) |',
        '| **Foundry 002** | Pipeline Builder for DD Data Ingestion | [bundle](articles/002_pipeline_builder_dd_data_ingestion/) |',
        '| **Foundry 003** | Workshop Application Patterns for Counterparty Risk Investigators | [starter bundle](articles/003_workshop_application_patterns_investigators/) |',
        '| **Foundry 004** | AIP-Driven Adverse-Media Summarization for DD Engagements | [starter bundle](articles/004_aip_adverse_media_summarization/) |',
        '| **Foundry 005** | Foundry Actions Framework for Audit-Trail Discipline | [starter bundle](articles/005_actions_framework_audit_trail_discipline/) |',
        '| **Foundry 006** | Code Repositories in Foundry | [starter bundle](articles/006_code_repositories_python_pyspark/) |',
        '| **Foundry 007** | Quiver for Ad-Hoc Counterparty Queries | [starter bundle](articles/007_quiver_ad_hoc_counterparty_queries/) |',
        '| **Foundry 008** | Time Series in Foundry | [starter bundle](articles/008_time_series_counterparty_risk_trajectories/) |',
        '| **Foundry 009** | Foundry Ontology Design for Beneficial-Ownership Networks | [starter bundle](articles/009_beneficial_ownership_networks/) |',
        '| **Foundry 010** | Foundry vs Snowflake + DBT vs Databricks for DD Analytics | [starter bundle](articles/010_platform_tradeoff_framework/) |',
        '| **Knowledge Graphs 001** | Beneficial-Ownership Lists to Cypher | [bundle](articles/knowledge_graphs_001_beneficial_ownership_cypher/) · [root csv](ownership_synthetic.csv) |',
        '| **Stochastic 002** | Modeling Journal-Entry Sequences in Production | [bundle](articles/stochastic_002_journal_entry_sequences/) · [script](stochastic_markov/companion_artifacts/002_journal_entry_production.py) |',
        '| **Stochastic 004** | Markov Mixture Models for Round-Tripping and Lapping Detection | [bundle](articles/stochastic_004_mixture_round_tripping/) · [script](stochastic_markov/companion_artifacts/004_mixture_round_tripping.py) |',
        '| **Stochastic 006** | Markov Decision Processes for Risk-Based Audit Sampling | [bundle](articles/stochastic_006_mdp_audit_sampling/) · [script](stochastic_markov/companion_artifacts/006_markov_decision_audit_sampling.py) |',
        '| **Stochastic 007** | Stochastic Volatility Models for Restatement-Timing Anomalies | [bundle](articles/stochastic_007_garch_restatement_timing/) · [script](stochastic_markov/companion_artifacts/007_garch_restatement_timing.py) · [notebook](notebooks/007_garch_restatement_timing.ipynb) |',
        '| **Stochastic 009** | Higher-Order and Variable-Order Markov Models for Long-Memory Fraud Schemes | [script](stochastic_markov/companion_artifacts/009_higher_order_variable_order_markov.py) |',
        '| **Stochastic 010** | Continuous-Time Markov Chains for Transaction Timing | [bundle](articles/stochastic_010_continuous_time_markov/) · [script](stochastic_markov/companion_artifacts/010_continuous_time_markov.py) |',
        '| **Excel 001** | Benford\'s Law in Excel | [bundle](articles/excel_001_benford_first_digit/) · [workbook](workbooks/001-benford-excel-template.xlsx) |',
        '| **Excel 002** | Same-Same-Different in Excel | [starter bundle](articles/excel_002_same_same_different/) · [workbook](workbooks/002-same-same-different.xlsx) |',
        '| **Excel 003** | Linear Regression Outlier Detection in Excel | [starter bundle](articles/excel_003_linear_regression_outlier_detection/) · [workbook](workbooks/003-linear-regression-outlier-detection.xlsx) |',
        '| **Excel 004** | Geospatial Address Mapping in Excel | [starter bundle](articles/excel_004_geospatial_address_mapping/) · [workbook](workbooks/004-geospatial-address-mapping.xlsx) |',
        '| **Excel 005** | OSINT for Financial Fraud in Excel | [bundle](articles/excel_005_osint_vendor_validation/) · [workbook](osint-vendor-validation.xlsx) · [cache template](engagement_files/osint_cache/%7BRequest_ID%7D.json) |',
        '| **Excel 006** | Time-Series Anomaly Detection in Excel | [bundle](articles/excel_006_time_series_anomaly_detection/) · [workbook](DD-Tech-006-TimeSeries.xlsx) |',
        '| **Excel 007** | Round-Number Bias and Threshold Avoidance in Excel | [starter bundle](articles/excel_007_round_number_bias_threshold_avoidance/) · [workbook](workbooks/007-round-number-bias-threshold-avoidance.xlsx) |',
        '| **Excel 008** | Date-Pattern Analysis | [bundle](articles/excel_008_date_pattern_analysis/) · [workbook](008_date_pattern_analysis.xlsx) |',
        '| **Excel 009** | Frequency-of-Amount Analysis | [bundle](articles/excel_009_frequency_of_amount/) · [workbook](DD-Tech-Lab-Repo/009-frequency-of-amount.xlsx) |',
        '| **Excel 010** | Correlation Diagnostics for Journal-Entry Pairs in Excel | [starter bundle](articles/excel_010_correlation_diagnostics_journal_entry_pairs/) · [workbook](workbooks/010-correlation-diagnostics-journal-entry-pairs.xlsx) |',
    ]
    INDEX.write_text('\n'.join(lines) + '\n')


def main():
    created_foundry = build_foundry()
    created_excel = build_excel()
    update_index()
    print(json.dumps({
        'created_foundry_bundles': created_foundry,
        'created_excel_bundles': created_excel,
        'count_foundry': len(created_foundry),
        'count_excel': len(created_excel),
    }, indent=2))


if __name__ == '__main__':
    main()
